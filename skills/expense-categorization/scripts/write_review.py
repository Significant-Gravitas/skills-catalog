#!/usr/bin/env python3
"""Validate a completed expense review table, split it into Ready to post and
Needs review, and tie its totals back to the source export.

    cd ~/skills/expense-categorization && python3 scripts/write_review.py \
        --review /home/user/work/review.csv --tx /home/user/work/amex.canonical.csv \
        [--chart chart.csv] [--control-total -900.69] \
        [--checks /home/user/work/checks.json] [--proposals /home/user/work/proposals.csv] \
        --outdir /home/user/out/expenses-2026-04

The review CSV uses templates/expense-review-template.csv. Checks, each fatal
(exit 2) because they mean the table no longer matches the source:
- every source row appears exactly once (source_file + source_row), or, when
  the owner split it, as parts numbered 1..n in `split_part` that keep the
  source currency and add up exactly to the source amount;
- every amount and currency equals the source row (no edits);
- every row has exactly the header's columns (a comma inside an unquoted
  field stops the run) and confidence is high, medium or low;
- every proposed_account is in --chart or is `unresolved`;
- an excluded row has an excluded_reason.
A row is Ready to post only if confidence is `high`, evidence is not
`missing`, proposed_account is not `unresolved`, review_need is empty and
special is empty or `none`. With `--checks checks.json` a row that
expense_checks.py lists as an exact duplicate or a special line is never
Ready; with `--proposals proposals.csv` a row whose confidence is above its
proposal cap is never Ready; a row whose receipt id also supports another row
is never Ready. Any other row goes to Needs review; demoted rows are listed
with the reason.

Writes ready-to-post.csv, needs-review.csv, totals.md (and review.xlsx with
three sheets if openpyxl is importable). Exit 1 when --control-total does not
tie. Never edits an input.
"""

import argparse
import os
import sys
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bkio  # noqa: E402

FIELDS = ["source_file", "source_row", "split_part", "txn_date", "posting_date", "vendor", "amount",
          "currency", "evidence", "payment_line", "proposed_account", "dimension", "reason", "confidence",
          "review_need", "basis", "special", "excluded_reason", "status"]
CONFIDENCE = ("high", "medium", "low")
RANK = {"low": 0, "medium": 1, "high": 2}
# first word of `evidence` that describes, rather than names, a document
NOT_AN_ID = {"", "missing", "receipt", "invoice", "contract", "none", "n/a", "-"}


def is_ready(r):
    def val(k, default=""):
        return (r.get(k) or default).strip().lower()
    return (val("confidence") == "high"
            and val("evidence", "missing") != "missing"
            and val("proposed_account", "unresolved") != "unresolved"
            and val("review_need") in ("", "-", "—")
            and val("special") in ("", "none")
            and val("excluded_reason") == "")


def load_checks(path):
    """Row ids that checks.json lists as an exact duplicate or a special line."""
    import json
    try:
        with open(path, encoding="utf-8") as fh:
            c = json.load(fh)
    except FileNotFoundError:
        raise bkio.InputError(f"--checks {path}: not found; run expense_checks.py --out first") from None
    except ValueError as e:
        raise bkio.InputError(f"--checks {path}: not valid JSON ({e})") from None
    flagged = {}
    for rid in c.get("exact_duplicate_rows") or []:
        flagged.setdefault(rid, []).append("exact duplicate in checks.json")
    for rid, kinds in (c.get("special_rows") or {}).items():
        flagged.setdefault(rid, []).append("special line " + "/".join(kinds) + " in checks.json")
    return flagged


def load_caps(path):
    """Row id -> confidence_cap from propose_from_history.py (blank = no cap)."""
    _, rows = bkio.read_csv(path)
    return {f"{r['source_file']}:{int(r['source_row'])}": (r.get("confidence_cap") or "").strip().lower()
            for r in rows}


def build(review_path, tx_path, chart_path=None, control_total=None, checks_path=None, proposals_path=None):
    tx = bkio.read_canonical(tx_path)
    src = {(t["source_file"], int(t["source_row"])): t for t in tx}
    header, rows = bkio.read_csv(review_path)
    missing_cols = [c for c in ("source_file", "source_row", "amount", "currency", "proposed_account",
                                "confidence", "evidence") if c not in header]
    if missing_cols:
        raise bkio.InputError(f"{review_path} lacks columns {missing_cols}; start from templates/expense-review-template.csv")
    chart = None
    if chart_path:
        _, crows = bkio.read_csv(chart_path)
        chart = {r["account"] for r in crows}
    flagged = load_checks(checks_path) if checks_path else {}
    caps = load_caps(proposals_path) if proposals_path else {}
    errors, seen, ready, review, demoted = [], set(), [], [], []
    parts = {}          # source key -> [(part number, amount, currency)]
    unsplit = set()     # source keys with an unsplit row
    receipt_rows = {}   # receipt id -> source keys it supports
    checked = []
    for r in rows:
        key = (r["source_file"], int(r["source_row"]))
        part = (r.get("split_part") or "").strip()
        t = src.get(key)
        if t is None:
            errors.append(f"row {key} is not in the source export")
            continue
        conf = (r.get("confidence") or "").strip().lower()
        if conf not in CONFIDENCE:
            errors.append(f"row {key}: confidence '{r.get('confidence')}' is not high, medium or low "
                          "(a shifted column? quote any field that contains a comma)")
        amt, _ = bkio.parse_amount(r["amount"])
        if part:
            if not part.isdigit() or int(part) < 1:
                errors.append(f"row {key}: split_part '{part}' must be 1, 2, 3 ...")
                continue
            parts.setdefault(key, []).append((int(part), amt, r["currency"] or ""))
        else:
            if key in seen:
                errors.append(f"row {key} appears twice (if the owner split it, number the parts in split_part)")
                continue
            unsplit.add(key)
            if amt != t["amount"] or (r["currency"] or "") != t["currency"]:
                errors.append(f"row {key}: amount {r['amount']} {r['currency']} differs from source "
                              f"{t['amount']} {t['currency']}")
        seen.add(key)
        acct = (r.get("proposed_account") or "").strip()
        if chart is not None and acct and acct.lower() != "unresolved" and acct not in chart:
            errors.append(f"row {key}: account '{acct}' is not in the chart")
        if (r.get("status") or "").lower() == "excluded" and not (r.get("excluded_reason") or "").strip():
            errors.append(f"row {key}: excluded without excluded_reason")
        ev = (r.get("evidence") or "").strip()
        doc = ev.split()[0].strip(",;:").lower() if ev else ""
        if doc not in NOT_AN_ID:
            receipt_rows.setdefault(doc, set()).add(key)
        checked.append((key, r))
    for key, plist in parts.items():
        t = src[key]
        nums = sorted(n for n, _, _ in plist)
        if key in unsplit:
            errors.append(f"row {key}: has both split parts and an unsplit row")
        if len(plist) < 2 or nums != list(range(1, len(plist) + 1)):
            errors.append(f"row {key}: split parts must be numbered 1..n with at least two parts, got {nums}")
        if any(c != t["currency"] for _, _, c in plist):
            errors.append(f"row {key}: every split part must keep the source currency {t['currency']}")
        total_parts = sum((a for _, a, _ in plist if a is not None), Decimal(0))
        if total_parts != t["amount"]:
            errors.append(f"row {key}: split parts add up to {total_parts}; the source is {t['amount']}")
    shared = {k for keys in receipt_rows.values() if len(keys) > 1 for k in keys}
    for key, r in checked:
        rid = f"{key[0]}:{key[1]}"
        reasons = list(flagged.get(rid, []))
        conf = (r.get("confidence") or "").strip().lower()
        cap = caps.get(rid, "")
        if cap in RANK and conf in RANK and RANK[conf] > RANK[cap]:
            reasons.append(f"confidence {conf} is above the proposal cap {cap}")
        if key in shared:
            reasons.append("its receipt also supports another source row")
        if is_ready(r) and not reasons:
            ready.append(r)
        else:
            if conf == "high":
                demoted.append(f"{rid}" + (f" ({'; '.join(reasons)})" if reasons else ""))
            review.append(r)
    for key in src:
        if key not in seen:
            errors.append(f"source row {key} is missing from the review")
    if errors:
        raise bkio.InputError("; ".join(errors[:20]) + (f" (+{len(errors) - 20} more)" if len(errors) > 20 else ""))
    total = sum((t["amount"] for t in tx), Decimal(0))
    ready_total = sum((bkio.parse_amount(r["amount"])[0] for r in ready), Decimal(0))
    review_total = sum((bkio.parse_amount(r["amount"])[0] for r in review), Decimal(0))
    unresolved = [r for r in review if (r.get("proposed_account") or "unresolved").lower() == "unresolved"]
    excluded = [r for r in rows if (r.get("status") or "").lower() == "excluded"]
    summary = {
        "rows": len(tx), "total": total, "ready_rows": len(ready), "ready_total": ready_total,
        "review_rows": len(review), "review_total": review_total, "unresolved_rows": len(unresolved),
        "excluded_rows": len(excluded), "demoted": demoted, "split_rows": sorted(parts),
        "ties_internal": ready_total + review_total == total,
        "control_total": control_total,
        "ties_control": None if control_total is None else total == control_total,
    }
    return ready, review, excluded, summary


def write_outputs(outdir, ready, review, excluded, s):
    os.makedirs(outdir, exist_ok=True)
    bkio.write_csv(os.path.join(outdir, "ready-to-post.csv"), FIELDS, ready)
    bkio.write_csv(os.path.join(outdir, "needs-review.csv"), FIELDS, review)
    lines = [
        "| Measure | Rows | Amount |", "| --- | --- | --- |",
        f"| Source export | {s['rows']} | {s['total']} |",
        f"| Ready to post | {s['ready_rows']} | {s['ready_total']} |",
        f"| Needs review | {s['review_rows']} | {s['review_total']} |",
        f"| of which unresolved | {s['unresolved_rows']} | |",
        f"| Excluded (listed with reason) | {s['excluded_rows']} | |",
        "",
        f"Ready + needs review = source: {'yes' if s['ties_internal'] else 'NO'}.",
    ]
    if s["control_total"] is not None:
        lines.append(f"Control total {s['control_total']}: {'ties' if s['ties_control'] else 'DOES NOT TIE'}.")
    if s["demoted"]:
        lines.append(f"Moved to Needs review (marked high, or blocked by checks, cap or shared receipt): {s['demoted']}")
    if s.get("split_rows"):
        lines.append(f"Split by the owner (parts add up to the source amount): {s['split_rows']}")
    lines.append("Script: write_review.py. Proposed categories are not posted entries and not tax advice.")
    with open(os.path.join(outdir, "totals.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    try:
        import openpyxl  # optional
        wb = openpyxl.Workbook()
        for title, data in (("Ready to post", ready), ("Needs review", review)):
            ws = wb.create_sheet(title)
            ws.append(FIELDS)
            for r in data:
                ws.append([r.get(f, "") for f in FIELDS])
        ws = wb.create_sheet("Totals")
        for line in lines:
            ws.append([line])
        wb.remove(wb["Sheet"])
        wb.save(os.path.join(outdir, "review.xlsx"))
    except ImportError:
        pass
    return "\n".join(lines)


def subprocess_run(argv):
    import subprocess
    return subprocess.run([sys.executable] + argv, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode


def selftest():
    import tempfile
    here = os.path.dirname(os.path.abspath(__file__))
    import normalize_export
    ex = os.path.join(here, "..", "examples", "bramble-april-card")
    rows = normalize_export.normalize(os.path.join(ex, "amex-2026-04.csv"), "%Y-%m-%d", flip=True)
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "amex.csv")
        bkio.write_csv(p, bkio.CANONICAL, rows)
        ready, review, excluded, s = build(os.path.join(ex, "expected-review.csv"), p,
                                           os.path.join(ex, "chart.csv"), Decimal("-900.69"))
        assert s["ties_internal"] and s["ties_control"], s
        assert [int(r["source_row"]) for r in ready] == [2], [r["source_row"] for r in ready]
        assert s["review_rows"] == 7, s
        # the fixture itself: header width on every row, confidence in range, caps respected
        import csv as _csv
        with open(os.path.join(ex, "expected-review.csv"), encoding="utf-8", newline="") as fh:
            raw = list(_csv.reader(fh))
        assert all(len(r) == len(raw[0]) for r in raw), [len(r) for r in raw]
        assert all(r[raw[0].index("confidence")] in CONFIDENCE for r in raw[1:])
        chk = os.path.join(d, "checks.json")
        rc = subprocess_run([os.path.join(here, "expense_checks.py"), "--tx", p, "--receipts",
                             os.path.join(ex, "receipts.csv"), "--out", chk])
        assert rc in (0, 1), rc
        prop = os.path.join(d, "proposals.csv")
        rc = subprocess_run([os.path.join(here, "propose_from_history.py"), "--tx", p,
                             "--rules", os.path.join(ex, "coding-rules.csv"),
                             "--history", os.path.join(ex, "history"), "--checks", chk,
                             "--chart", os.path.join(ex, "chart.csv"), "--out", prop])
        assert rc == 0, rc
        _, _, _, s2 = build(os.path.join(ex, "expected-review.csv"), p, os.path.join(ex, "chart.csv"),
                            None, chk, prop)
        assert s2["ready_rows"] == 1, s2
        assert not any("above the proposal cap" in x for x in s2["demoted"]), s2["demoted"]
        # a duplicate or refund marked high is not Ready once checks.json is given
        with open(os.path.join(ex, "expected-review.csv"), encoding="utf-8", newline="") as fh:
            table = list(_csv.DictReader(fh))
        for r in table:
            if r["source_row"] in ("7", "8"):
                r.update(confidence="high", evidence="R9 receipt (ties)" if r["source_row"] == "7" else "R8 receipt (ties)",
                         proposed_account="Meals and subsistence" if r["source_row"] == "7" else "Software subscriptions",
                         review_need="", special="none")
        tweak = os.path.join(d, "tweak.csv")
        bkio.write_csv(tweak, FIELDS, table)
        r_plain = build(tweak, p, os.path.join(ex, "chart.csv"))[3]
        assert r_plain["ready_rows"] == 3, r_plain
        r_checked = build(tweak, p, os.path.join(ex, "chart.csv"), None, chk, prop)[3]
        assert r_checked["ready_rows"] == 1, r_checked["demoted"]
        # one receipt supporting two rows is not Ready
        for r in table:
            if r["source_row"] == "7":
                r["evidence"] = "R1 receipt (ties)"
        bkio.write_csv(tweak, FIELDS, table)
        assert build(tweak, p)[3]["ready_rows"] == 1, build(tweak, p)[3]
        # an owner split: parts numbered 1..n, same currency, adding up exactly
        split = [dict(r) for r in table if r["source_row"] != "4"]
        amz = next(r for r in table if r["source_row"] == "4")
        split += [dict(amz, split_part="1", amount="-600.00", proposed_account="Office supplies"),
                  dict(amz, split_part="2", amount="-12.99", proposed_account="Office supplies")]
        bkio.write_csv(tweak, FIELDS, split)
        s3 = build(tweak, p, os.path.join(ex, "chart.csv"), Decimal("-900.69"))[3]
        assert s3["ties_control"] and s3["split_rows"] == [("amex-2026-04.csv", 4)], s3
        split[-1]["amount"] = "-13.00"
        bkio.write_csv(tweak, FIELDS, split)
        try:
            build(tweak, p)
            raise AssertionError("split parts that do not add up must fail")
        except bkio.InputError as e:
            assert "split parts add up to" in str(e), e
        write_outputs(d, ready, review, excluded, s)
        assert os.path.exists(os.path.join(d, "totals.md"))
        # tampering with an amount must fail
        bad = os.path.join(d, "bad.csv")
        with open(os.path.join(ex, "expected-review.csv"), encoding="utf-8") as fh:
            text = fh.read().replace("-45.00", "-54.00", 1)
        with open(bad, "w", encoding="utf-8") as fh:
            fh.write(text)
        try:
            build(bad, p)
            raise AssertionError("edited amount must fail")
        except bkio.InputError as e:
            assert "differs from source" in str(e)
    print("selftest ok")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--review")
    ap.add_argument("--tx")
    ap.add_argument("--chart")
    ap.add_argument("--control-total")
    ap.add_argument("--checks")
    ap.add_argument("--proposals")
    ap.add_argument("--outdir")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if not (a.review and a.tx and a.outdir):
        bkio.die("--review, --tx and --outdir are required")
    try:
        ct = bkio.parse_amount(a.control_total)[0] if a.control_total else None
        ready, review, excluded, s = build(a.review, a.tx, a.chart, ct, a.checks, a.proposals)
    except bkio.InputError as e:
        bkio.die(str(e), "correct the review table (never the source export) and rerun")
    print(write_outputs(a.outdir, ready, review, excluded, s))
    if s["ties_control"] is False:
        sys.exit(1)


if __name__ == "__main__":
    main()
