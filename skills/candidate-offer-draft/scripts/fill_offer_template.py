#!/usr/bin/env python3
"""Fill the company's offer template from approved terms, leaving every other word untouched.

Usage:
  cd ~/skills/candidate-offer-draft && \
  python3 scripts/fill_offer_template.py <template.docx|template.md|template.txt> <terms.json> <out-file> [--date-style long|iso]

  .docx needs python-docx:  python3 -c "import docx" || pip install --user python-docx

Placeholders in the template are {{key}} (spaces inside the braces allowed), where key
is a term in terms.json ("base_pay", "start_date", ...) or candidate.<field>
("candidate.legal_name", "candidate.chosen_name").

Rules
  * Only text inside {{...}} changes. Clause wording, order, headers, footers
    and tables are left as they are. In .docx, a placeholder split across
    Word runs is still found; the replacement keeps the first run's formatting.
  * A term with no value, or a value without both a source and an approver,
    is written as "[OWNER TO CONFIRM: key]" (highlighted yellow in .docx) and
    reported. It is never guessed.
  * List values (e.g. contingencies) are joined with "; ".
  * start_date and response_by (YYYY-MM-DD in terms.json) are written out as
    "6 January 2027" (--date-style long, the default). Use --date-style iso
    only if the company's template requires ISO dates.
  * The output file name must contain "DRAFT"; if it does not, "DRAFT-" is
    prefixed.
Report on stdout: filled, open (OWNER TO CONFIRM), refused (no source or
approver), template placeholders with no term, and approved terms the template
never uses (a clause may be missing from the template; tell the owner, and
do not add wording).
Exit 0 = all filled, 1 = written with open items, 2 = bad input (including a
term that is not an object: FORMAT).
"""

import copy
import json
import re
import sys
from datetime import date
from pathlib import Path

DATE_TERMS = ("start_date", "response_by")
PH = re.compile(r"\{\{\s*([A-Za-z0-9_.]+)\s*\}\}")


def as_text(v) -> str:
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    return "" if v is None else str(v)


class Resolver:
    def __init__(self, data: dict, date_style: str = "long"):
        self.date_style = date_style
        self.terms = data.get("terms") or {}
        self.cand = data.get("candidate") or {}
        self.filled, self.open, self.refused, self.unknown = set(), set(), set(), set()

    def __call__(self, key: str) -> tuple[str, bool]:
        if key.startswith("candidate."):
            v = as_text(self.cand.get(key.split(".", 1)[1])).strip()
            if v:
                self.filled.add(key)
                return v, False
            self.open.add(key)
            return f"[OWNER TO CONFIRM: {key}]", True
        t = self.terms.get(key)
        if t is None:
            self.unknown.add(key)
            return f"[OWNER TO CONFIRM: {key}]", True
        v = as_text(t.get("value")).strip()
        if not v:
            self.open.add(key)
            return f"[OWNER TO CONFIRM: {key}]", True
        if not as_text(t.get("source")).strip() or not as_text(t.get("approver")).strip():
            self.refused.add(key)
            return f"[OWNER TO CONFIRM: {key}]", True
        self.filled.add(key)
        if key in DATE_TERMS and self.date_style == "long":
            try:
                d = date.fromisoformat(v)
                v = f"{d.day} {d:%B %Y}"
            except ValueError:
                pass
        return v, False


def fill_text(text: str, resolve: Resolver) -> str:
    return PH.sub(lambda m: resolve(m.group(1))[0], text)


def docx_paragraphs(doc):
    def walk(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                     s.even_page_header, s.even_page_footer):
            try:
                yield from walk(part)
            except Exception:
                continue


def fill_paragraph(p, resolve: Resolver, Run, YELLOW) -> None:
    text = "".join(r.text for r in p.runs)
    matches = list(PH.finditer(text))
    for m in reversed(matches):
        value, missing = resolve(m.group(1))
        s, e = m.span()
        runs = p.runs
        bounds, pos = [], 0
        for r in runs:
            bounds.append((pos, pos + len(r.text)))
            pos += len(r.text)
        i = next(k for k, (a, b) in enumerate(bounds) if a <= s < b)
        j = next(k for k, (a, b) in enumerate(bounds) if a < e <= b)
        ri, rj = runs[i], runs[j]
        before = ri.text[: s - bounds[i][0]]
        after = rj.text[e - bounds[j][0]:]
        for k in range(i + 1, j):
            runs[k].text = ""
        new_el = copy.deepcopy(ri._r)
        ri._r.addnext(new_el)
        new_run = Run(new_el, p)
        new_run.text = value
        if missing:
            new_run.font.highlight_color = YELLOW
        if i == j:
            tail_el = copy.deepcopy(ri._r)
            new_el.addnext(tail_el)
            Run(tail_el, p).text = after
            ri.text = before
        else:
            ri.text = before
            rj.text = after


def main() -> int:
    args = sys.argv[1:]
    date_style = "long"
    if "--date-style" in args:
        i = args.index("--date-style")
        if i + 1 >= len(args) or args[i + 1] not in ("long", "iso"):
            print("error: --date-style must be long or iso", file=sys.stderr)
            return 2
        date_style = args[i + 1]
        del args[i:i + 2]
    if len(args) != 3:
        print(__doc__.split("Placeholders")[0], file=sys.stderr)
        return 2
    tpl, terms_path, out = (Path(x).expanduser() for x in args)
    if not tpl.is_file():
        print(f"error: template not found: {tpl}", file=sys.stderr)
        return 2
    try:
        data = json.loads(terms_path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        print(f"error: cannot read terms: {exc}", file=sys.stderr)
        return 2
    if not isinstance(data, dict) or not isinstance(data.get("terms") or {}, dict) \
            or not isinstance(data.get("candidate") or {}, dict):
        print("error: terms.json must be an object with \"candidate\" and \"terms\" objects "
              "(see templates/offer-terms.json)", file=sys.stderr)
        return 2
    bad = [k for k, t in (data.get("terms") or {}).items() if t is not None and not isinstance(t, dict)]
    if bad:
        print(f"error: FORMAT: term(s) {bad} must each be an object with value, source, approver, "
              "approved_on; run check_terms.py and fix terms.json first", file=sys.stderr)
        return 2
    if "DRAFT" not in out.name.upper():
        out = out.with_name("DRAFT-" + out.name)
    out.parent.mkdir(parents=True, exist_ok=True)
    resolve = Resolver(data, date_style)

    if tpl.suffix.lower() == ".docx":
        try:
            import docx
            from docx.enum.text import WD_COLOR_INDEX
            from docx.text.run import Run
        except ImportError:
            print("error: python-docx missing: pip install --user python-docx", file=sys.stderr)
            return 2
        doc = docx.Document(str(tpl))
        seen = 0
        for p in docx_paragraphs(doc):
            if PH.search("".join(r.text for r in p.runs)):
                seen += 1
                fill_paragraph(p, resolve, Run, WD_COLOR_INDEX.YELLOW)
        leftover = [p.text for p in docx_paragraphs(doc) if "{{" in p.text]
        doc.save(str(out))
        if leftover:
            print(f"WARNING placeholder-like text left unfilled (malformed braces?): {leftover[:3]}")
    elif tpl.suffix.lower() in (".md", ".txt"):
        out.write_text(fill_text(tpl.read_text(encoding="utf-8-sig"), resolve), encoding="utf-8")
    else:
        print("error: template must be .docx, .md or .txt; ask the owner for one of those", file=sys.stderr)
        return 2

    used = resolve.filled | resolve.open | resolve.refused
    unused = sorted(k for k, t in (data.get("terms") or {}).items()
                    if k not in used and as_text((t or {}).get("value")).strip())
    print(f"wrote {out}")
    print(f"filled:  {sorted(resolve.filled)}")
    print(f"open [OWNER TO CONFIRM]: {sorted(resolve.open)}")
    print(f"refused (no source or approver): {sorted(resolve.refused)}")
    print(f"template placeholders with no term: {sorted(resolve.unknown)}")
    print(f"approved terms the template never uses: {unused}")
    return 1 if (resolve.open or resolve.refused or resolve.unknown) else 0


if __name__ == "__main__":
    sys.exit(main())
