#!/usr/bin/env python3
r"""Validate offer terms before any draft is shown: sources, approvers, dates, pay history, and email/letter consistency.

Usage:
  cd ~/skills/candidate-offer-draft && \
  python3 scripts/check_terms.py <terms.json> [--email email.txt] [--letter letter.md|letter.docx] \
      [--today YYYY-MM-DD] [--json]

terms.json (see templates/offer-terms.json):
  {"candidate": {"chosen_name": ..., "legal_name": ..., "candidate_id": ...},
   "jurisdiction": "<hiring jurisdiction from the hiring plan, e.g. US-CA, UK>",
   "fcra_disclosure_on_record": true|false|null,   (for a non-US hire: the
                            candidate's background-check notice/consent is on record)
   "terms": {"<key>": {"value": ..., "source": ..., "approver": ..., "approved_on": "YYYY-MM-DD"}, ...}}
  A term with an empty value is OPEN and is shown as [OWNER TO CONFIRM].

ERROR (exit 1):
  NO_SOURCE / NO_APPROVER   a filled term without its source or approver
  PAY_HISTORY               any term, source or note that refers to current, previous
                            or expected-from-history pay (salary-history bans, e.g.
                            Cal. Labor Code 432.3; other states in HR Dive's list [56][75])
  INFERRED                  a source that is market data, a benchmark, an estimate,
                            another candidate or an earlier draft
  DATE_FORMAT               start_date, response_by or approved_on not YYYY-MM-DD
  RESPONSE_PAST             response_by is before --today
  START_BEFORE_RESPONSE     start_date is before response_by
  CURRENCY / PAY_PERIOD     base_pay present without a 3-letter currency or a pay period
  CONTINGENCY_SOURCE        a contingency whose source is not the company's offer template
  FORMAT                    a term that is not an object with value/source/approver/approved_on
  EMAIL_MISMATCH / LETTER_MISMATCH  an amount or date in the email/letter that is not in the terms,
                            including bare amounts ("65,000", "salary of 65000"), trailing
                            symbols and grouped thousands ("65.000 EUR", "65 000 GBP"),
                            and percentages ("a 10% bonus") that no term states.
                            Approved amounts come only from money-bearing terms: base_pay,
                            and amounts written in other terms such as bonus_text or a
                            relocation term. Numbers inside dates, level, location,
                            work_terms, contingencies and other non-money terms never
                            count as approved amounts. An amount shown as the pay rate
                            (a pay word such as salary/pay/rate/wage before it and a pay
                            period after it) must equal base_pay.
  CURRENCY_MISMATCH         an amount in the email/letter in a currency other than terms.currency
                            (a bare "$" or the word "dollars" counts as USD only when
                            terms.currency is USD; "pounds"/"sterling" = GBP, "euros" = EUR,
                            "yen" = JPY when written after the amount)
  PAY_PERIOD_MISMATCH       the base-pay amount shown with a period ("per month", "annual")
                            other than terms.pay_period
  AMBIGUOUS_DATE            a numeric date such as 06/01/2027 in the email or letter
                            A date without a year ("Friday 16 October") whose day and
                            month match no approved date is an EMAIL/LETTER_MISMATCH.
  SPELLED_AMOUNT            an amount written in words ("sixty-five thousand pounds") with
                            no amount in figures on the same line: it cannot be checked;
                            write it in figures with the approved currency
WARNING (exit 0 unless errors):
  OPEN                      a required term without a value ([OWNER TO CONFIRM])
  FCRA_DISCLOSURE           a background-check contingency and no standalone FCRA
                            disclosure recorded before the report is procured [48]
                            (US hires, or when no jurisdiction is recorded)
  LOCAL_CHECK_RULES         the same for a non-US jurisdiction: US FCRA wording does
                            not apply; local rules may (e.g. UK DBS, UK GDPR Art. 10)
  NO_CURRENCY               a bare amount in the email/letter with no currency shown
  ISO_DATE_IN_TEXT          a YYYY-MM-DD date in the email/letter; write "6 January 2027"
  YEARLESS_DATE             a date without a year whose day and month match an approved
                            date; add the year so it cannot be misread
  SPELLED_AMOUNT            an amount in words next to one in figures on the same line:
                            confirm the words say the same approved amount
  MEDICAL                   a medical-exam contingency: post-offer only, and only if
                            every entrant in the job category gets it [46]
  APPROVED_IN_FUTURE        approved_on is after --today

Exit 0 = no errors, 1 = errors, 2 = bad input (unreadable JSON, or a top level
or "terms" that is not an object).
"""

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

REQUIRED = ["entity", "title", "level", "manager", "location", "work_terms", "start_date",
            "base_pay", "currency", "pay_period", "bonus_text", "equity_text", "benefits_source",
            "contingencies", "response_by", "signatory"]
DATE_KEYS = ["start_date", "response_by"]
PERIODS = {"year", "annual", "annum", "month", "week", "day", "hour", "per year", "per month",
           "per week", "per hour", "per annum"}
PAY_HISTORY = re.compile(r"(current|previous|prior|last|past|existing|present)\s+(salary|pay|comp\w*|base|package|earnings)|"
                         r"salary history|pay history|what (she|he|they) (makes|made|earns|earned)|"
                         r"match(ing)? (her|his|their) (salary|pay|offer)|\+\s?\d+\s?% (on|over|above) (her|his|their)", re.I)
INFERRED = re.compile(r"market (data|rate)|benchmark|glassdoor|levels\.fyi|payscale|estimat|typical|"
                      r"usual|standard (rate|package)|similar (candidate|hire|role)|another candidate|"
                      r"last hire|previous (offer|draft)|earlier draft|same as (the )?last", re.I)
# A number next to a currency may use "." or space thousands grouping
# ("65.000 EUR", "65 000 GBP"); exactly-three-digit groups are read as thousands.
GROUPED = r"\d{1,3}(?:[. \u00a0]\d{3})+(?![\d,])"
NUMBER = r"(?:" + GROUPED + r"|\d[\d,]*(?:\.\d+)?)"
MONEY = re.compile(r"(?:(?P<sym>[£$€¥])\s?(?P<a>\d[\d,]*(?:\.\d+)?)\s?(?P<k>[kK])?)|"
                   r"(?:(?P<c1>[A-Z]{3})\s?(?P<b>" + NUMBER + r")\s?(?P<k2>[kK])?)|"
                   r"(?:(?P<c>" + NUMBER + r")\s?(?P<k3>[kK])?\s?(?P<c2>GBP|USD|EUR|CAD|AUD|INR|JPY|CHF|SEK|NOK|DKK|NZD|SGD))|"
                   r"(?:(?P<d>" + NUMBER + r")\s?(?P<k4>[kK])?\s?(?P<sym2>[£$€¥]))")
PERCENT = re.compile(r"(?<![\d.,])(\d+(?:\.\d+)?)\s?(?:%|per ?cent\b)", re.I)
# Terms whose numbers are never amounts (dates, level, section numbers, days on site).
NON_MONEY_KEYS = set(["entity", "title", "level", "manager", "location", "work_terms", "start_date",
                      "response_by", "currency", "pay_period", "benefits_source", "contingencies",
                      "signatory"])
PAY_WORD_BEFORE = re.compile(r"\b(salary|pay|rate|wage|base)\b(?![^.;\n]*\b(bonus|allowance|relocation|sign-?on|"
                             r"stipend|expenses?|review)\b)[^.;\n]{0,30}$", re.I)
NUMERIC_DATE = re.compile(r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b")
MONTHS = "january february march april may june july august september october november december".split()
WORD_DATE = re.compile(r"\b(\d{1,2})(?:st|nd|rd|th)?\s+(" + "|".join(m[:3] + r"[a-z]*" for m in MONTHS) + r")\s+(\d{4})\b|"
                       r"\b(" + "|".join(m[:3] + r"[a-z]*" for m in MONTHS) + r")\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})\b", re.I)
ISO_DATE = re.compile(r"\b(\d{4})-(\d{2})-(\d{2})\b")
# Dates with no year. Month names are exact (full or 3-4 letter abbreviation)
# so words such as "Marketing" are not read as months; the month-first form
# needs a capitalised month so the verb "may" is not read as May.
_MON = r"(january|february|march|april|may|june|july|august|september|october|november|december|" \
       r"(?:jan|feb|mar|apr|jun|jul|aug|sept?|oct|nov|dec)\.?)"
_MON_CAP = r"(?-i:(January|February|March|April|May|June|July|August|September|October|November|December|" \
           r"Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sept?|Oct|Nov|Dec))\.?"
YEARLESS_DATE = re.compile(r"\b(\d{1,2})(?:st|nd|rd|th)?\s+(?:of\s+)?" + _MON + r"(?![a-z])|"
                           r"\b" + _MON_CAP + r"\s+(\d{1,2})(?:st|nd|rd|th)?\b(?!\s*,?\s*\d{4})(?![\d,.]*\d)", re.I)
# Currency written as a word after an amount ("62,000 euros").
CCY_WORD = re.compile(r"^\s*(?:[kK]\s*)?(pounds?(?:\s+sterling)?|sterling|quid|euros?|dollars?|yen)\b", re.I)
CCY_WORD_MAP = {"pound": "GBP", "sterling": "GBP", "quid": "GBP", "euro": "EUR", "dollar": "$", "yen": "JPY"}
# Amounts written in words.
_NUMWORD = r"(one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|thirteen|fourteen|fifteen|" \
           r"sixteen|seventeen|eighteen|nineteen|twenty|thirty|forty|fifty|sixty|seventy|eighty|ninety|hundred|a)"
SPELLED = re.compile(r"\b" + _NUMWORD + r"(?:[- ]" + _NUMWORD + r")*[- ](thousand|million|hundred)\b", re.I)
# Bare amounts: digits with thousands separators, or any number right after a pay word.
BARE_MONEY = re.compile(r"(?<![\w£$€¥.,])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?\s?[kK]\b)(?![\w.,]*\d)|"
                        r"(?:salary|pay|base|wage|compensation|bonus|rate)\s+(?:of\s+|is\s+|will be\s+)?(\d[\d,]*(?:\.\d+)?)(?!\s?[£$€¥]|\s?(?-i:[A-Z]{3})\b)",
                        re.I)
SYMBOL_CCY = {"£": "GBP", "€": "EUR", "¥": "JPY"}
PERIOD_AFTER = re.compile(r"^[^.;\n]{0,30}?\b(?:(?:per|a|an|each|every)\s+(year|annum|month|week|day|hour)|"
                          r"(annual(?:ly)?|yearly|monthly|weekly|daily|hourly|p\.?\s?a\.?(?=\W|$)))", re.I)
PERIOD_BEFORE = re.compile(r"\b(annual|yearly|monthly|weekly|daily|hourly)\s+(?:base\s+)?(?:salary|pay|rate|wage)\b[^.;\n]{0,20}$", re.I)
PERIOD_CANON = {"year": "year", "annum": "year", "annual": "year", "annually": "year", "yearly": "year",
                "pa": "year", "p.a.": "year", "p.a": "year", "month": "month", "monthly": "month",
                "week": "week", "weekly": "week", "day": "day", "daily": "day", "hour": "hour", "hourly": "hour"}


def as_text(v) -> str:
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    return "" if v is None else str(v)


def iso(v: str):
    try:
        return date.fromisoformat(v)
    except (TypeError, ValueError):
        return None


def num(s: str, k: bool) -> float:
    s = s.strip()
    if re.fullmatch(GROUPED, s):
        n = float(re.sub(r"[. \u00a0]", "", s))
    else:
        n = float(s.replace(",", ""))
    return n * 1000 if k else n


def money_in(text: str) -> list[tuple[str, float]]:
    return [(tok, n) for tok, n, _, _, _ in money_detail(text)]


def canon_period(word: str) -> str:
    w = word.lower().replace(" ", "")
    return PERIOD_CANON.get(w, PERIOD_CANON.get(w.rstrip("."), w))


def money_detail(text: str) -> list[tuple[str, float, str, int, int]]:
    """(token, amount, currency or '' or '$', start, end) for every amount, including bare ones."""
    out, spans = [], []
    for m in MONEY.finditer(text):
        if m.group("a"):
            sym = m.group("sym")
            out.append((m.group(0), num(m.group("a"), bool(m.group("k"))), SYMBOL_CCY.get(sym, sym), m.start(), m.end()))
        elif m.group("b"):
            out.append((m.group(0), num(m.group("b"), bool(m.group("k2"))), m.group("c1"), m.start(), m.end()))
        elif m.group("c"):
            out.append((m.group(0), num(m.group("c"), bool(m.group("k3"))), m.group("c2"), m.start(), m.end()))
        elif m.group("d"):
            sym = m.group("sym2")
            out.append((m.group(0), num(m.group("d"), bool(m.group("k4"))), SYMBOL_CCY.get(sym, sym), m.start(), m.end()))
        else:
            continue
        spans.append((m.start(), m.end()))
    for m in BARE_MONEY.finditer(text):
        g = 1 if m.group(1) else 2
        s, e = m.start(g), m.end(g)
        if any(a <= s < b or a < e <= b for a, b in spans):
            continue
        raw = m.group(g)
        k = raw[-1:].lower() == "k"
        try:
            n = num(raw.rstrip("kK").strip(), k)
        except ValueError:
            continue
        cw = CCY_WORD.match(text[e:e + 25])
        cur = ""
        if cw:
            word = cw.group(1).lower().split()[0].rstrip("s")
            cur = CCY_WORD_MAP.get(word, CCY_WORD_MAP.get(cw.group(1).lower(), ""))
        out.append((raw, n, cur, s, e))
    return out


def period_near(text: str, start: int, end: int) -> str:
    m = PERIOD_AFTER.search(text[end:end + 40])
    if m:
        return canon_period(m.group(1) or m.group(2))
    m = PERIOD_BEFORE.search(text[max(0, start - 40):start])
    if m:
        return canon_period(m.group(1))
    return ""


def dates_in(text: str) -> list[tuple[str, date | None]]:
    out = []
    for m in WORD_DATE.finditer(text):
        if m.group(1):
            d, mon, y = m.group(1), m.group(2), m.group(3)
        else:
            mon, d, y = m.group(4), m.group(5), m.group(6)
        idx = [x[:3] for x in MONTHS].index(mon[:3].lower()) + 1
        try:
            out.append((m.group(0), date(int(y), idx, int(d))))
        except ValueError:
            out.append((m.group(0), None))
    for m in ISO_DATE.finditer(text):
        out.append((m.group(0), iso(m.group(0))))
    return out


def yearless_in(text: str) -> list[tuple[str, int, int]]:
    """(token, month, day) for dates written without a year, outside full dates."""
    taken = [(m.start(), m.end()) for m in WORD_DATE.finditer(text)]
    out = []
    for m in YEARLESS_DATE.finditer(text):
        if any(a <= m.start() < b for a, b in taken):
            continue
        if m.group(1):
            d, mon = m.group(1), m.group(2)
        else:
            mon, d = m.group(3), m.group(4)
        idx = [x[:3] for x in MONTHS].index(mon[:3].lower()) + 1
        out.append((m.group(0), idx, int(d)))
    return out


def read_doc(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            print("error: reading .docx needs python-docx: pip install --user python-docx", file=sys.stderr)
            sys.exit(2)
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    return path.read_text(encoding="utf-8-sig")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("terms")
    ap.add_argument("--email")
    ap.add_argument("--letter")
    ap.add_argument("--today", default=date.today().isoformat())
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    try:
        data = json.loads(Path(a.terms).expanduser().read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        print(f"error: cannot read terms: {exc}", file=sys.stderr)
        return 2
    today = iso(a.today)
    if not today:
        print("error: --today must be YYYY-MM-DD", file=sys.stderr)
        return 2
    if not isinstance(data, dict) or not isinstance(data.get("terms") or {}, dict):
        print("error: terms.json must be an object with a \"terms\" object (see templates/offer-terms.json)",
              file=sys.stderr)
        return 2
    F = []

    def add(level, rule, key, msg):
        F.append({"level": level, "rule": rule, "term": key, "message": msg})

    terms = {}
    for key, t in (data.get("terms") or {}).items():
        if isinstance(t, dict):
            terms[key] = t
        else:
            add("ERROR", "FORMAT", key, "each term must be an object with value, source, approver, approved_on "
                                        f"(got {type(t).__name__}); treated as open")
            terms[key] = {}

    for key in REQUIRED:
        t = terms.get(key) or {}
        if not as_text(t.get("value")).strip():
            add("WARNING", "OPEN", key, "no approved value: [OWNER TO CONFIRM]")
    for key, t in terms.items():
        val = as_text(t.get("value")).strip()
        blob = " ".join(as_text(t.get(f)) for f in ("value", "source", "note", "approver"))
        if PAY_HISTORY.search(blob):
            add("ERROR", "PAY_HISTORY", key, f"refers to pay history ('{PAY_HISTORY.search(blob).group(0)}'); never use or derive terms from it")
        if val and INFERRED.search(as_text(t.get("source"))):
            add("ERROR", "INFERRED", key, f"source '{t.get('source')}' is not an approval; terms come only from an authorised person or the template")
        if val and not as_text(t.get("source")).strip():
            add("ERROR", "NO_SOURCE", key, "filled without a source")
        if val and not as_text(t.get("approver")).strip():
            add("ERROR", "NO_APPROVER", key, "filled without an approver")
        if t.get("approved_on"):
            d = iso(t["approved_on"])
            if not d:
                add("ERROR", "DATE_FORMAT", key, f"approved_on '{t['approved_on']}' is not YYYY-MM-DD")
            elif d > today:
                add("WARNING", "APPROVED_IN_FUTURE", key, f"approved_on {d} is after today {today}")
    for key in DATE_KEYS:
        v = as_text((terms.get(key) or {}).get("value")).strip()
        if v and not iso(v):
            add("ERROR", "DATE_FORMAT", key, f"'{v}' is not YYYY-MM-DD")
    rb = iso(as_text((terms.get("response_by") or {}).get("value")))
    sd = iso(as_text((terms.get("start_date") or {}).get("value")))
    if rb and rb < today:
        add("ERROR", "RESPONSE_PAST", "response_by", f"{rb} is before today {today}")
    if rb and sd and sd < rb:
        add("ERROR", "START_BEFORE_RESPONSE", "start_date", f"start {sd} is before response-by {rb}")
    if as_text((terms.get("base_pay") or {}).get("value")).strip():
        cur = as_text((terms.get("currency") or {}).get("value")).strip()
        per = as_text((terms.get("pay_period") or {}).get("value")).strip().lower()
        if not re.fullmatch(r"[A-Z]{3}", cur):
            add("ERROR", "CURRENCY", "currency", f"base pay needs a 3-letter currency code, got '{cur}'")
        if per not in PERIODS:
            add("ERROR", "PAY_PERIOD", "pay_period", f"base pay needs a pay period ({', '.join(sorted(PERIODS))}), got '{per}'")
    cont = terms.get("contingencies") or {}
    items = cont.get("value") if isinstance(cont.get("value"), list) else ([cont["value"]] if cont.get("value") else [])
    if items and "template" not in as_text(cont.get("source")).lower():
        add("ERROR", "CONTINGENCY_SOURCE", "contingencies", "contingency lines come only from the company's offer template")
    joined = " ".join(str(x) for x in items).lower()
    juris = as_text(data.get("jurisdiction")).strip()
    is_us = (not juris) or bool(re.search(r"\b(US|USA|U\.S\.(A\.)?|United States)\b", juris))
    if re.search(r"background|criminal|credit|consumer report|reference check", joined) and data.get("fcra_disclosure_on_record") is not True:
        if is_us:
            add("WARNING", "FCRA_DISCLOSURE", "contingencies",
                "background-check contingency: confirm a standalone written disclosure and authorisation were (or will be) obtained before the report is procured [48]"
                + ("" if juris else " (no hiring jurisdiction recorded; add it from the hiring plan)"))
        else:
            add("WARNING", "LOCAL_CHECK_RULES", "contingencies",
                f"background-check contingency, jurisdiction '{juris}': US FCRA wording does not apply; local background-check "
                "rules may (e.g. UK DBS, UK GDPR Art. 10). Confirm the notice/consent step with the people lead")
    if re.search(r"medical|health|drug|physical exam", joined):
        add("WARNING", "MEDICAL", "contingencies",
            "medical exam or inquiry: only after the offer, and only if every entrant in the job category gets it [46]")

    known_money, known_dates, known_pct = set(), set(), set()
    for key, t in terms.items():
        text = as_text(t.get("value"))
        d = iso(text.strip())
        if d:
            known_dates.add(d)
        known_pct.update(float(x) for x in PERCENT.findall(text))
        if key == "base_pay":
            for tok in re.findall(r"\d[\d,]*(?:\.\d+)?", text):
                known_money.add(float(tok.replace(",", "")))
        elif key not in NON_MONEY_KEYS:
            for _, n in money_in(ISO_DATE.sub(" ", text)):
                known_money.add(n)
    for flag, rule in ((a.email, "EMAIL_MISMATCH"), (a.letter, "LETTER_MISMATCH")):
        if not flag:
            continue
        p = Path(flag).expanduser()
        if not p.is_file():
            print(f"error: not found: {p}", file=sys.stderr)
            return 2
        text = read_doc(p)
        ccy = as_text((terms.get("currency") or {}).get("value")).strip().upper()
        per = canon_period(as_text((terms.get("pay_period") or {}).get("value")).strip()) if \
            as_text((terms.get("pay_period") or {}).get("value")).strip() else ""
        base_vals = {n for _, n in money_in(as_text((terms.get("base_pay") or {}).get("value")))}
        for tok in re.findall(r"\d[\d,]*(?:\.\d+)?", as_text((terms.get("base_pay") or {}).get("value"))):
            base_vals.add(float(tok.replace(",", "")))
        for tok, n, cur, s, e in money_detail(text):
            tok = tok.strip()
            if n not in known_money:
                add("ERROR", rule, p.name, f"amount '{tok}' is not in the approved terms")
            if cur == "$":
                if ccy != "USD":
                    add("ERROR", "CURRENCY_MISMATCH", p.name,
                        f"amount '{tok}' uses '$' but the approved currency is '{ccy or '?'}'; write the approved currency")
            elif cur and ccy and cur.upper() != ccy:
                add("ERROR", "CURRENCY_MISMATCH", p.name, f"amount '{tok}' is in {cur} but the approved currency is {ccy}")
            elif not cur:
                add("WARNING", "NO_CURRENCY", p.name, f"amount '{tok}' has no currency; show {ccy or 'the approved currency'}")
            if n in base_vals and per:
                seen = period_near(text, s, e)
                if seen and seen != per:
                    add("ERROR", "PAY_PERIOD_MISMATCH", p.name,
                        f"amount '{tok}' is shown per {seen} but the approved pay period is {per}")
            elif n in known_money and base_vals and period_near(text, s, e) and \
                    PAY_WORD_BEFORE.search(text[max(0, s - 40):s]):
                add("ERROR", rule, p.name,
                    f"amount '{tok}' is shown as the pay rate, but the approved base pay is "
                    f"{', '.join(f'{v:g}' for v in sorted(base_vals))}")
        for m in PERCENT.finditer(text):
            if float(m.group(1)) not in known_pct:
                add("ERROR", rule, p.name, f"percentage '{m.group(0)}' is not in the approved terms "
                                           "(bonus_text, equity_text or another approved term)")
        for m in ISO_DATE.finditer(text):
            add("WARNING", "ISO_DATE_IN_TEXT", p.name, f"'{m.group(0)}': write dates out, e.g. '6 January 2027'")
        for tok, d in dates_in(text):
            if d is None or d not in known_dates:
                add("ERROR", rule, p.name, f"date '{tok}' is not an approved term date")
        for tok, mon, day in yearless_in(text):
            if any((k.month, k.day) == (mon, day) for k in known_dates):
                add("WARNING", "YEARLESS_DATE", p.name, f"date '{tok}' has no year; write it in full, e.g. '6 January 2027'")
            else:
                add("ERROR", rule, p.name, f"date '{tok}' (no year) matches no approved term date")
        for line in text.splitlines():
            for m in SPELLED.finditer(line):
                if money_detail(line):
                    add("WARNING", "SPELLED_AMOUNT", p.name,
                        f"'{m.group(0)}' is written in words; confirm it says the same approved amount as the figures")
                else:
                    add("ERROR", "SPELLED_AMOUNT", p.name,
                        f"'{m.group(0)}' is an amount in words the check cannot verify; write it in figures with the approved currency")
        for m in NUMERIC_DATE.finditer(text):
            add("ERROR", "AMBIGUOUS_DATE", p.name, f"'{m.group(0)}' can be read two ways; write '6 January 2027'")
        if PAY_HISTORY.search(text):
            add("ERROR", "PAY_HISTORY", p.name, f"mentions pay history ('{PAY_HISTORY.search(text).group(0)}')")

    errors = [f for f in F if f["level"] == "ERROR"]
    if a.json:
        print(json.dumps(F, indent=2))
    else:
        for f in sorted(F, key=lambda x: x["level"] != "ERROR"):
            print(f"{f['level']:7} {f['rule']:22} {f['term']}: {f['message']}")
        print(f"{len(errors)} error(s), {len(F) - len(errors)} warning(s).")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
